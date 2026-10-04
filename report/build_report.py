"""Builds the interim report (Part B) as .docx from results/summary.json and figures, then converts to PDF."""
import json
import os
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import content as C

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RESULTS = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "results")
FIG = os.path.join(RESULTS, "figures")
OUT = os.path.join(sys.argv[2] if len(sys.argv) > 2 else HERE, "Interim_Report.docx")
FONT = "Times New Roman"
HEADER_FILL = "DCE6F2"


class Report:
    def __init__(self):
        self.doc = Document()
        self.fig_no = 0
        self.tab_no = 0
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.1

    def para(self, text="", bold=False, italic=False, size=None, align=None, after=None, before=None):
        p = self.doc.add_paragraph()
        self._runs(p, text, bold, italic, size)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if before is not None:
            p.paragraph_format.space_before = Pt(before)
        return p

    @staticmethod
    def _runs(p, text, bold=False, italic=False, size=None):
        """Supports **bold** spans inside text."""
        for i, chunk in enumerate(text.split("**")):
            if not chunk:
                continue
            r = p.add_run(chunk)
            r.bold = bold or i % 2 == 1
            r.italic = italic
            r.font.name = FONT
            if size:
                r.font.size = Pt(size)

    def body(self, text):
        return self.para(text, align="justify")

    def h1(self, text):
        p = self.para(text, bold=True, size=13, before=10, after=4)
        p.paragraph_format.keep_with_next = True

    def h2(self, text):
        p = self.para(text, bold=True, size=11.5, before=6, after=3)
        p.paragraph_format.keep_with_next = True

    def bullets(self, items, style="List Bullet"):
        for it in items:
            p = self.doc.add_paragraph(style=style)
            self._runs(p, it)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(2)

    def caption(self, text, kind):
        if kind == "Figure":
            self.fig_no += 1
            label = f"Figure {self.fig_no}: "
        else:
            self.tab_no += 1
            label = f"Table {self.tab_no}: "
        p = self.para("", align="center", after=8 if kind == "Figure" else 3)
        r = p.add_run(label)
        r.bold, r.font.size, r.font.name = True, Pt(9.5), FONT
        r = p.add_run(text)
        r.italic, r.font.size, r.font.name = True, Pt(9.5), FONT
        if kind == "Table":
            p.paragraph_format.keep_with_next = True

    def figure(self, name, caption, width=16.0):
        path = os.path.join(FIG, name)
        if not os.path.exists(path):
            self.para(f"[missing figure {name}]", italic=True)
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Cm(width))
        self.caption(caption, "Figure")

    def table(self, header, rows, caption=None, widths=None, size=9, align_cols=None):
        if caption:
            self.caption(caption, "Table")
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(header):
            self._cell(t.cell(0, j), h, bold=True, size=size, fill=HEADER_FILL, center=True)
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row):
                center = align_cols is not None and j in align_cols
                self._cell(t.cell(i, j), str(v), size=size, center=center)
        if widths:
            t.autofit = False
            for j, w in enumerate(widths):
                t.columns[j].width = Cm(w)
            for row in t.rows:
                for j, w in enumerate(widths):
                    row.cells[j].width = Cm(w)
        self._repeat_header(t)
        self.para("", after=4)
        return t

    def _cell(self, cell, text, bold=False, size=9, fill=None, center=False):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._runs(p, text, bold=bold, size=size)
        if fill:
            tc = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fill)
            tc.append(shd)

    @staticmethod
    def _repeat_header(t):
        tr = t.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        tr.append(el)

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Cm(0.3)
            r = p.add_run(line if line else " ")
            r.font.name = "Courier New"
            r.font.size = Pt(8)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            ppr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "F3F3F1")
            ppr.append(shd)
        self.para("", after=2)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def page_numbers(self):
        for sec in self.doc.sections:
            p = sec.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run()
            for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
                if tag:
                    el = OxmlElement("w:fldChar")
                    el.set(qn("w:fldCharType"), tag)
                else:
                    el = OxmlElement("w:instrText")
                    el.set(qn("xml:space"), "preserve")
                    el.text = txt
                r._r.append(el)
            r.font.size = Pt(9)

    def save(self, path):
        self.page_numbers()
        self.doc.save(path)


def cover(R):
    R.para("", after=30)
    for line in ("ADAPTIVE DEEP NEURAL NETWORK", "INFERENCE", "USING CONFIDENCE-BASED EARLY EXITS"):
        R.para(line, bold=True, size=18, align="center", after=0)
    R.para("", after=18)
    for line in ("Interim Report on", "Deep Learning Project", "[ICT-4442]"):
        R.para(line, bold=True, size=12, align="center", after=2)
    R.para("", after=14)
    R.para("Submitted By", size=10.5, align="center", after=10)
    for name, reg in C.TEAM:
        R.para(f"{name}  |  {reg}", bold=True, size=10, align="center", after=1)
    R.para("", after=16)
    logo = os.path.join(HERE, "assets", "mahe_logo.png")
    if os.path.exists(logo):
        p = R.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo, width=Cm(10.5))
    R.para("", after=16)
    for line in ("SCHOOL OF COMPUTER ENGINEERING", "MANIPAL INSTITUTE OF TECHNOLOGY", "MANIPAL ACADEMY OF HIGHER EDUCATION"):
        R.para(line, bold=True, size=10, align="center", after=1)
    R.para("", after=10)
    R.para(C.MONTH, bold=True, size=10, align="center")
    R.page_break()


def main():
    with open(os.path.join(RESULTS, "summary.json")) as f:
        S = json.load(f)
    R = Report()
    cover(R)
    C.write_body(R, S)
    R.save(OUT)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(OUT), OUT], check=True,
                   stdout=subprocess.DEVNULL)
    print("wrote", OUT, "and", OUT.replace(".docx", ".pdf"))


if __name__ == "__main__":
    main()
